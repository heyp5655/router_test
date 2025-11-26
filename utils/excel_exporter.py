"""Excel导出工具类"""
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any


class ExcelExporter:
    """测试结果Excel导出器"""

    def __init__(self, output_dir: str = "reports"):
        """
        初始化Excel导出器

        Args:
            output_dir: 输出目录
        """
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    def export_test_results(self, results: List[Dict[str, Any]],
                           router_config: Dict[str, Any],
                           test_info: Dict[str, Any]) -> str:
        """
        导出测试结果到Excel

        Args:
            results: 测试结果列表
            router_config: 路由器配置信息
            test_info: 测试信息（包含总体结果、耗时等）

        Returns:
            str: 生成的Excel文件路径
        """
        # 创建工作簿
        wb = Workbook()
        ws = wb.active
        ws.title = "测试结果"

        # 设置列宽
        ws.column_dimensions['A'].width = 15  # 标签列
        ws.column_dimensions['B'].width = 35  # 测试项
        ws.column_dimensions['C'].width = 40  # 测试用例名称
        ws.column_dimensions['D'].width = 12  # 测试结果
        ws.column_dimensions['E'].width = 60  # 备注

        # 定义样式
        header_font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        # 边框样式
        thin_border = Border(
            left=Side(style='thin', color='000000'),
            right=Side(style='thin', color='000000'),
            top=Side(style='thin', color='000000'),
            bottom=Side(style='thin', color='000000')
        )

        # 添加汇总信息（在表格之前）
        current_row = 1
        # "测试汇总"标题 - 居中对齐
        title_cell = ws.cell(row=current_row, column=1)
        title_cell.value = "测试汇总"
        title_cell.font = Font(name='微软雅黑', size=12, bold=True)
        title_cell.alignment = Alignment(horizontal='center', vertical='center')  # 居中对齐
        ws.merge_cells(f'A{current_row}:E{current_row}')
        current_row += 1

        summary_data = [
            ("路由器型号:", router_config.get('model', '')),
            ("测试时间:", datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
            ("总用例数:", test_info.get('total', 0)),
            ("通过数:", test_info.get('passed', 0)),
            ("失败数:", test_info.get('failed', 0)),
            ("错误数:", test_info.get('error', 0)),
            ("总耗时:", f"{test_info.get('duration', 0)}秒"),
            ("总体结果:", test_info.get('overall_result', 'UNKNOWN'))
        ]

        # 汇总信息 - 靠左对齐
        left_alignment_summary = Alignment(horizontal='left', vertical='center')
        for label, value in summary_data:
            # 标签列（靠左）
            label_cell = ws.cell(row=current_row, column=1)
            label_cell.value = label
            label_cell.font = Font(name='微软雅黑', size=10, bold=True)
            label_cell.alignment = left_alignment_summary

            # 值列（靠左）
            value_cell = ws.cell(row=current_row, column=2)
            value_cell.value = value
            value_cell.font = Font(name='微软雅黑', size=10)
            value_cell.alignment = left_alignment_summary

            current_row += 1

        # 空一行
        current_row += 1

        # 设置表头
        header_row = current_row
        headers = ['ID', '测试项', '测试用例名称', '测试结果', '备注']
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col_num)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # PASS和FAIL的样式
        pass_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
        pass_font = Font(name='微软雅黑', size=10, color='006100')

        fail_fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
        fail_font = Font(name='微软雅黑', size=10, color='9C0006')

        error_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
        error_font = Font(name='微软雅黑', size=10, color='9C5700')

        normal_font = Font(name='微软雅黑', size=10)
        center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

        # 填充数据（从表头下一行开始）
        data_start_row = header_row + 1
        for idx, result in enumerate(results):
            row_num = data_start_row + idx

            # ID
            cell = ws.cell(row=row_num, column=1)
            cell.value = idx + 1
            cell.font = normal_font
            cell.alignment = center_alignment
            cell.border = thin_border

            # 测试项
            cell = ws.cell(row=row_num, column=2)
            cell.value = result.get('category', '')
            cell.font = normal_font
            cell.alignment = left_alignment
            cell.border = thin_border

            # 测试用例名称
            cell = ws.cell(row=row_num, column=3)
            cell.value = result.get('test_name', '')
            cell.font = normal_font
            cell.alignment = left_alignment
            cell.border = thin_border

            # 测试结果
            status = result.get('status', '')
            cell = ws.cell(row=row_num, column=4)
            cell.value = status
            cell.alignment = center_alignment
            cell.border = thin_border

            # 根据测试结果设置颜色
            if status == 'PASS':
                cell.fill = pass_fill
                cell.font = pass_font
            elif status == 'FAIL':
                cell.fill = fail_fill
                cell.font = fail_font
            elif status == 'ERROR':
                cell.fill = error_fill
                cell.font = error_font
            else:
                cell.font = normal_font

            # 备注（如果失败，填写失败原因）
            cell = ws.cell(row=row_num, column=5)
            if status in ['FAIL', 'ERROR']:
                message = result.get('message', '')
                cell.value = message
            else:
                cell.value = ''
            cell.font = normal_font
            cell.alignment = left_alignment
            cell.border = thin_border

        # 生成文件名
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"测试报告_{timestamp}.xlsx"
        filepath = os.path.join(self.output_dir, filename)

        # 保存文件
        wb.save(filepath)
        print(f"Excel报告已生成: {filepath}")

        return filepath
