#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""读取用例步骤.xlsx中的用例ID 11"""

import openpyxl
import sys

# 设置输出编码为UTF-8
sys.stdout.reconfigure(encoding='utf-8')

# 读取Excel文件
wb = openpyxl.load_workbook(r'E:\GIT\ROUTER_TEST\用例步骤.xlsx')
ws = wb.active

# 打印表头
print("=" * 100)
print("表头信息：")
headers = []
for cell in ws[1]:
    if cell.value:
        headers.append(cell.value)
        print(f"  列{cell.column}: {cell.value}")
print("=" * 100)

# 查找用例ID 11
print("\n查找用例ID 11...")
found = False
for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
    if row[0].value == 11 or str(row[0].value) == '11':
        found = True
        print("\n" + "=" * 100)
        print(f"✅ 找到用例ID 11 (第{row_idx}行)")
        print("=" * 100)

        for i, cell in enumerate(row):
            if cell.value is not None and str(cell.value).strip():
                header = headers[i] if i < len(headers) else f"列{i}"
                print(f"\n【{header}】:")
                print(f"{cell.value}")
                print("-" * 100)
        break

if not found:
    print("\n❌ 未找到用例ID 11")
    print("\n可用的用例ID列表：")
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            print(f"  - ID {row[0]}: {row[1] if len(row) > 1 else '(无名称)'}")

wb.close()
