"""
路由器页面结构化分析工具 V5 - 按功能区域组织

核心改进：
1. ✅ 识别每个页面的功能区域（Section）
2. ✅ 按层级组织元素：主菜单 > 子菜单 > 功能区域 > 配置项
3. ✅ 生成易于配置的Excel表格
4. ✅ 每个配置项清楚标明所属功能区域
5. ✅ 方便用户直接填写配置值

使用方法：
python scripts/analyze_router_pages_v5.py --ip 192.168.3.1 --username admin --password admin1
"""

import sys
import os
import time
import argparse
from datetime import datetime
import json
import re

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from core.router_client import RouterClient


class RouterPageAnalyzerV5:
    """路由器页面结构化分析器 V5"""

    def __init__(self, router_ip: str, username: str, password: str, model: str = "UR35"):
        self.router_ip = router_ip
        self.username = username
        self.password = password
        self.model = model

        self.router_client = RouterClient(
            router_ip=router_ip,
            username=username,
            password=password,
            model=model
        )

        self.menu_items = []
        self.structured_data = []  # 结构化数据：按页面和功能区域组织

        self.screenshot_dir = r"E:\GIT\ROUTER_TEST\logs\screenshots"
        os.makedirs(self.screenshot_dir, exist_ok=True)

    def analyze(self):
        """执行完整的分析流程"""
        print("\n" + "=" * 80)
        print("路由器页面结构化分析工具 V5")
        print("=" * 80)
        print(f"路由器IP: {self.router_ip}")
        print(f"用户名: {self.username}")
        print(f"型号: {self.model}")
        print("=" * 80 + "\n")

        try:
            # 1. 登录
            print("步骤1: 登录路由器...")
            if not self.router_client.login_web():
                raise Exception("登录失败")
            print("✅ 登录成功\n")

            # 2. 识别菜单
            print("步骤2: 识别菜单结构...")
            self._collect_all_menu_items()
            print(f"✅ 发现 {len(self.menu_items)} 个页面\n")

            # 3. 结构化分析每个页面
            if len(self.menu_items) > 0:
                print("步骤3: 结构化分析页面（按功能区域组织）...")
                self._analyze_pages_structured()
                print(f"✅ 分析完成\n")

            # 4. 生成友好的Excel
            print("步骤4: 生成结构化Excel配置表...")
            output_file = self._generate_friendly_excel()
            print(f"✅ Excel已生成: {output_file}\n")

            # 5. 统计
            self._print_statistics()

            print("=" * 80)
            print("分析完成！")
            print("=" * 80)

        except Exception as e:
            print(f"\n❌ 分析失败: {e}")
            import traceback
            traceback.print_exc()
        finally:
            if self.router_client.driver:
                self.router_client.driver.quit()

    def _collect_all_menu_items(self):
        """收集所有菜单项"""
        driver = self.router_client.driver
        try:
            time.sleep(3)
            main_menu = driver.find_element(By.ID, "mainmenu")
            main_menu_items = main_menu.find_elements(By.XPATH, ".//li[contains(@class, 'dropdown')]")
            print(f"  发现 {len(main_menu_items)} 个主菜单项")

            for idx, main_item in enumerate(main_menu_items, 1):
                try:
                    main_link = main_item.find_element(By.XPATH, "./a")
                    main_text = main_link.text.strip()
                    main_target = main_link.get_attribute('data-target')

                    print(f"\n  [{idx}] {main_text}")

                    if main_target:
                        self.menu_items.append({
                            'main_menu': main_text,
                            'sub_menu': '',
                            'text': main_text,
                            'target': main_target
                        })

                    try:
                        if 'dropdown-toggle' in main_link.get_attribute('class'):
                            main_link.click()
                            time.sleep(0.5)

                        sub_menu = main_item.find_element(By.XPATH, ".//ul[contains(@class, 'dropdown-menu')]")
                        sub_items = sub_menu.find_elements(By.XPATH, ".//li/a")

                        for sub_item in sub_items:
                            try:
                                sub_text = sub_item.text.strip()
                                sub_target = sub_item.get_attribute('data-target')
                                if sub_text and sub_target:
                                    self.menu_items.append({
                                        'main_menu': main_text,
                                        'sub_menu': sub_text,
                                        'text': f"{main_text} > {sub_text}",
                                        'target': sub_target
                                    })
                                    print(f"      • {sub_text}")
                            except:
                                continue
                    except:
                        pass
                except Exception as e:
                    continue
        except Exception as e:
            print(f"  ❌ 菜单收集失败: {e}")

    def _analyze_pages_structured(self):
        """结构化分析每个页面"""
        driver = self.router_client.driver
        total = len(self.menu_items)

        for idx, menu_item in enumerate(self.menu_items, 1):
            print(f"\n  [{idx}/{total}] 分析: {menu_item['text']}")

            try:
                # 导航
                url = f"http://{self.router_ip}/#{menu_item['target']}"
                driver.get(url)
                time.sleep(5)  # 增加等待时间，确保动态内容加载

                # 额外等待：等待主内容区域出现
                try:
                    WebDriverWait(driver, 10).until(
                        EC.presence_of_element_located((By.ID, "maincontent"))
                    )
                    time.sleep(2)  # 再等2秒让内容完全渲染
                except:
                    pass

                # 截图
                safe_filename = menu_item['text'].replace('>', '_').replace('/', '_')[:50]
                screenshot = self._take_screenshot(f"v5_page_{idx}_{safe_filename}")

                # 结构化分析页面
                page_structure = self._analyze_page_structure(menu_item, screenshot)

                self.structured_data.append(page_structure)

                # 统计
                section_count = len(page_structure['sections'])
                config_count = sum(len(s['config_items']) for s in page_structure['sections'])
                button_count = sum(len(s['buttons']) for s in page_structure['sections'])

                print(f"      ✅ {section_count}个功能区域, {config_count}个配置项, {button_count}个按钮")

            except Exception as e:
                print(f"      ⚠️  分析出错: {e}")
                continue

    def _analyze_page_structure(self, menu_item, screenshot):
        """分析单个页面的结构"""
        driver = self.router_client.driver

        page_structure = {
            'main_menu': menu_item['main_menu'],
            'sub_menu': menu_item['sub_menu'],
            'page_name': menu_item['text'],
            'target': menu_item['target'],
            'screenshot': screenshot,
            'sections': []
        }

        # 使用JavaScript分析页面结构
        structure_script = """
        function analyzePage() {
            var result = {
                sections: [],
                topLevelButtons: []
            };

            // 1. 查找所有ys-section（功能区域）
            var sections = document.querySelectorAll('.ys-section');

            // 如果没有找到ys-section，创建一个默认section（整个页面）
            if (sections.length === 0) {
                var mainContent = document.querySelector('#maincontent');
                if (mainContent) {
                    sections = [mainContent];
                }
            }

            for (var i = 0; i < sections.length; i++) {
                var section = sections[i];
                var sectionInfo = {
                    sectionTitle: '',
                    sectionId: section.id || '',
                    configItems: [],
                    buttons: []
                };

                // 获取section标题
                var titleElem = section.querySelector('.ys-section-title, h3.ys-title-3, h3');
                if (titleElem) {
                    sectionInfo.sectionTitle = titleElem.textContent.trim();
                }

                // 如果section是maincontent且没有标题，使用默认标题
                if (!sectionInfo.sectionTitle && section.id === 'maincontent') {
                    sectionInfo.sectionTitle = '主配置区域';
                }

                // 在这个section内查找表单元素
                var inputs = section.querySelectorAll('input:not([type="hidden"])');
                var selects = section.querySelectorAll('select');
                var textareas = section.querySelectorAll('textarea');

                // 处理input
                for (var j = 0; j < inputs.length; j++) {
                    var input = inputs[j];
                    if (!input.offsetParent) continue; // 跳过隐藏元素

                    var item = {
                        type: 'input',
                        inputType: input.type || 'text',
                        id: input.id || '',
                        name: input.name || '',
                        value: input.value || '',
                        placeholder: input.placeholder || '',
                        checked: input.checked || false,
                        disabled: input.disabled || false,
                        required: input.required || false,
                        labelText: '',
                        xpath: ''
                    };

                    // 查找label
                    if (input.id) {
                        var label = section.querySelector('label[for="' + input.id + '"]');
                        if (label) {
                            item.labelText = label.textContent.trim();
                        }
                    }

                    // 如果没找到label，查找父元素中的ys-label
                    if (!item.labelText) {
                        var parent = input.closest('.row');
                        if (parent) {
                            var ysLabel = parent.querySelector('.ys-label');
                            if (ysLabel) {
                                item.labelText = ysLabel.textContent.trim();
                            }
                        }
                    }

                    // 生成XPath
                    if (input.id) {
                        item.xpath = '//*[@id="' + input.id + '"]';
                    }

                    sectionInfo.configItems.push(item);
                }

                // 处理select
                for (var j = 0; j < selects.length; j++) {
                    var select = selects[j];
                    if (!select.offsetParent) continue;

                    var options = [];
                    for (var k = 0; k < select.options.length; k++) {
                        options.push({
                            value: select.options[k].value,
                            text: select.options[k].text
                        });
                    }

                    var item = {
                        type: 'select',
                        id: select.id || '',
                        name: select.name || '',
                        value: select.value || '',
                        disabled: select.disabled || false,
                        labelText: '',
                        options: options,
                        xpath: ''
                    };

                    // 查找label
                    if (select.id) {
                        var label = section.querySelector('label[for="' + select.id + '"]');
                        if (label) {
                            item.labelText = label.textContent.trim();
                        }
                    }

                    if (!item.labelText) {
                        var parent = select.closest('.row');
                        if (parent) {
                            var ysLabel = parent.querySelector('.ys-label');
                            if (ysLabel) {
                                item.labelText = ysLabel.textContent.trim();
                            }
                        }
                    }

                    if (select.id) {
                        item.xpath = '//*[@id="' + select.id + '"]';
                    }

                    sectionInfo.configItems.push(item);
                }

                // 处理textarea
                for (var j = 0; j < textareas.length; j++) {
                    var textarea = textareas[j];
                    if (!textarea.offsetParent) continue;

                    var item = {
                        type: 'textarea',
                        id: textarea.id || '',
                        name: textarea.name || '',
                        value: textarea.value || '',
                        placeholder: textarea.placeholder || '',
                        disabled: textarea.disabled || false,
                        labelText: '',
                        xpath: ''
                    };

                    if (textarea.id) {
                        var label = section.querySelector('label[for="' + textarea.id + '"]');
                        if (label) {
                            item.labelText = label.textContent.trim();
                        }
                    }

                    if (!item.labelText) {
                        var parent = textarea.closest('.row');
                        if (parent) {
                            var ysLabel = parent.querySelector('.ys-label');
                            if (ysLabel) {
                                item.labelText = ysLabel.textContent.trim();
                            }
                        }
                    }

                    if (textarea.id) {
                        item.xpath = '//*[@id="' + textarea.id + '"]';
                    }

                    sectionInfo.configItems.push(item);
                }

                // 查找section内的按钮
                var buttons = section.querySelectorAll('button');
                for (var j = 0; j < buttons.length; j++) {
                    var btn = buttons[j];
                    if (!btn.offsetParent) continue;

                    sectionInfo.buttons.push({
                        text: btn.textContent.trim(),
                        id: btn.id || '',
                        className: btn.className || '',
                        type: btn.type || 'button'
                    });
                }

                // 只添加有内容的section
                if (sectionInfo.configItems.length > 0 || sectionInfo.buttons.length > 0) {
                    result.sections.push(sectionInfo);
                }
            }

            // 2. 查找顶层按钮（不在section内的）
            var topButtons = document.querySelectorAll('#maincontent > button, #maincontent .button, button.button');
            for (var i = 0; i < topButtons.length; i++) {
                var btn = topButtons[i];
                if (!btn.offsetParent) continue;

                // 检查是否已在section内
                var inSection = btn.closest('.ys-section');
                if (!inSection) {
                    result.topLevelButtons.push({
                        text: btn.textContent.trim(),
                        id: btn.id || '',
                        className: btn.className || ''
                    });
                }
            }

            return result;
        }

        return analyzePage();
        """

        try:
            js_result = driver.execute_script(structure_script)

            # 转换JavaScript结果到Python
            for section in js_result.get('sections', []):
                section_data = {
                    'section_title': section.get('sectionTitle', '(未命名区域)'),
                    'section_id': section.get('sectionId', ''),
                    'config_items': [],
                    'buttons': []
                }

                # 配置项
                for item in section.get('configItems', []):
                    config_item = {
                        'field_type': item.get('type', ''),
                        'input_type': item.get('inputType', ''),
                        'label': item.get('labelText', '(无标签)'),
                        'id': item.get('id', ''),
                        'name': item.get('name', ''),
                        'current_value': item.get('value', ''),
                        'placeholder': item.get('placeholder', ''),
                        'xpath': item.get('xpath', ''),
                        'disabled': item.get('disabled', False),
                        'required': item.get('required', False),
                        'checked': item.get('checked', False),
                        'options': json.dumps(item.get('options', []), ensure_ascii=False) if item.get('options') else ''
                    }
                    section_data['config_items'].append(config_item)

                # 按钮
                for btn in section.get('buttons', []):
                    section_data['buttons'].append({
                        'text': btn.get('text', ''),
                        'id': btn.get('id', ''),
                        'type': btn.get('type', 'button')
                    })

                page_structure['sections'].append(section_data)

            # 顶层按钮
            page_structure['top_buttons'] = []
            for btn in js_result.get('topLevelButtons', []):
                page_structure['top_buttons'].append({
                    'text': btn.get('text', ''),
                    'id': btn.get('id', '')
                })

        except Exception as e:
            print(f"        JavaScript分析出错: {e}")

        return page_structure

    def _take_screenshot(self, name):
        """截图"""
        try:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"{name}_{timestamp}.png"
            filepath = os.path.join(self.screenshot_dir, filename)
            self.router_client.driver.save_screenshot(filepath)
            return filename
        except:
            return ''

    def _generate_friendly_excel(self):
        """生成友好的Excel配置表"""
        wb = Workbook()
        if 'Sheet' in wb.sheetnames:
            del wb['Sheet']

        # Sheet1: 配置总览
        ws_overview = wb.create_sheet("配置总览", 0)
        self._create_overview_sheet(ws_overview)

        # Sheet2: 按页面分组的配置表（主要的配置表）
        ws_config = wb.create_sheet("路由器配置表", 1)
        self._create_friendly_config_sheet(ws_config)

        # Sheet3: 页面结构说明
        ws_structure = wb.create_sheet("页面结构", 2)
        self._create_structure_sheet(ws_structure)

        # 保存
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_dir = r"E:\GIT\ROUTER_TEST\config"
        os.makedirs(output_dir, exist_ok=True)
        output_file = os.path.join(output_dir, f"router_config_template_v5_{timestamp}.xlsx")
        wb.save(output_file)

        return output_file

    def _create_overview_sheet(self, ws):
        """创建配置总览表"""
        ws.append(['路由器配置总览'])
        ws.append([''])
        ws.append(['主菜单', '子菜单', '功能区域数', '配置项数', '按钮数'])

        # 样式
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        for cell in ws[3]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # 数据
        for page in self.structured_data:
            section_count = len(page['sections'])
            config_count = sum(len(s['config_items']) for s in page['sections'])
            button_count = sum(len(s['buttons']) for s in page['sections'])

            ws.append([
                page['main_menu'],
                page['sub_menu'],
                section_count,
                config_count,
                button_count
            ])

        # 列宽
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 15

        ws.merge_cells('A1:E1')
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')

    def _create_friendly_config_sheet(self, ws):
        """创建友好的配置表（按页面和功能区域组织）"""
        # 说明行
        ws.append(['路由器配置参数表（按页面功能区域组织）'])
        ws.append(['说明：在"配置值"列填写您要配置的值，"是否启用"设为Y或N'])
        ws.append([''])

        # 表头
        headers = ['主菜单', '子菜单', '功能区域', '配置项名称', '字段类型', '当前值',
                   '配置值', '是否启用', '备注（选项/说明）', '定位ID', 'XPath']
        ws.append(headers)

        # 表头样式
        header_fill = PatternFill(start_color='FFC000', end_color='FFC000', fill_type='solid')
        header_font = Font(bold=True, size=11)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for cell in ws[4]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border

        # 数据行
        row_num = 5
        current_main = None
        current_sub = None

        for page in self.structured_data:
            main_menu = page['main_menu']
            sub_menu = page['sub_menu']
            page_has_data = False

            for section in page['sections']:
                if not section['config_items']:
                    continue

                page_has_data = True
                section_title = section['section_title']

                for item in section['config_items']:
                    # 字段类型说明
                    field_type_desc = item['field_type']
                    if item['field_type'] == 'input':
                        field_type_desc = f"输入框({item['input_type']})"
                    elif item['field_type'] == 'select':
                        field_type_desc = "下拉选择"
                    elif item['field_type'] == 'textarea':
                        field_type_desc = "文本域"

                    # 备注（选项列表）
                    remark = ''
                    if item['options']:
                        try:
                            options = json.loads(item['options'])
                            remark = ' | '.join([f"{o['text']}({o['value']})" for o in options[:5]])
                            if len(options) > 5:
                                remark += f" ...共{len(options)}项"
                        except:
                            pass

                    if item['placeholder']:
                        remark = f"提示: {item['placeholder']}" + (' | ' + remark if remark else '')

                    row = [
                        main_menu,
                        sub_menu,
                        section_title,
                        item['label'],
                        field_type_desc,
                        item['current_value'][:30] if item['current_value'] else '',
                        '',  # 配置值（用户填写）
                        'Y' if not item['disabled'] else 'N',  # 是否启用
                        remark,
                        item['id'],
                        item['xpath'][:50] if item['xpath'] else ''
                    ]
                    ws.append(row)

                    # 样式
                    for cell in ws[row_num]:
                        cell.border = border
                        cell.alignment = Alignment(vertical='center', wrap_text=True)

                    # 配置值列（G列）加高亮
                    ws[f'G{row_num}'].fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')

                    row_num += 1

            # 页面之间加空行
            if page_has_data:
                ws.append([''] * 11)
                row_num += 1

        # 列宽
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 25
        ws.column_dimensions['E'].width = 15
        ws.column_dimensions['F'].width = 20
        ws.column_dimensions['G'].width = 25
        ws.column_dimensions['H'].width = 10
        ws.column_dimensions['I'].width = 40
        ws.column_dimensions['J'].width = 20
        ws.column_dimensions['K'].width = 40

        # 合并说明
        ws.merge_cells('A1:K1')
        ws.merge_cells('A2:K2')
        ws['A1'].font = Font(bold=True, size=14, color='FF0000')
        ws['A1'].alignment = Alignment(horizontal='center')
        ws['A2'].font = Font(size=11, color='0000FF')
        ws['A2'].alignment = Alignment(horizontal='center')

        # 冻结窗格（冻结表头）
        ws.freeze_panes = 'A5'

    def _create_structure_sheet(self, ws):
        """创建页面结构说明表"""
        ws.append(['页面结构说明'])
        ws.append([''])
        ws.append(['主菜单', '子菜单', '功能区域', '配置项', '按钮'])

        header_fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        for cell in ws[3]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for page in self.structured_data:
            for section in page['sections']:
                config_items = ', '.join([item['label'] for item in section['config_items'][:5]])
                if len(section['config_items']) > 5:
                    config_items += f" ...共{len(section['config_items'])}项"

                buttons = ', '.join([btn['text'] for btn in section['buttons']])

                ws.append([
                    page['main_menu'],
                    page['sub_menu'],
                    section['section_title'],
                    config_items,
                    buttons
                ])

        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 30
        ws.column_dimensions['D'].width = 50
        ws.column_dimensions['E'].width = 30

        ws.merge_cells('A1:E1')
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center')

    def _print_statistics(self):
        """打印统计信息"""
        total_sections = sum(len(p['sections']) for p in self.structured_data)
        total_configs = sum(
            sum(len(s['config_items']) for s in p['sections'])
            for p in self.structured_data
        )
        total_buttons = sum(
            sum(len(s['buttons']) for s in p['sections'])
            for p in self.structured_data
        )

        print("\n" + "=" * 80)
        print("统计信息")
        print("=" * 80)
        print(f"总页面数: {len(self.structured_data)}")
        print(f"总功能区域: {total_sections}")
        print(f"总配置项: {total_configs}")
        print(f"总按钮数: {total_buttons}")
        print("=" * 80)


def main():
    parser = argparse.ArgumentParser(description='路由器页面结构化分析工具 V5')
    parser.add_argument('--ip', required=True, help='路由器IP地址')
    parser.add_argument('--username', default='admin', help='用户名')
    parser.add_argument('--password', default='password', help='密码')
    parser.add_argument('--model', default='UR35', help='型号')

    args = parser.parse_args()

    analyzer = RouterPageAnalyzerV5(
        router_ip=args.ip,
        username=args.username,
        password=args.password,
        model=args.model
    )

    analyzer.analyze()


if __name__ == '__main__':
    main()
