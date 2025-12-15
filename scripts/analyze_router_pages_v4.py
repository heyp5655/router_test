"""
路由器页面完整元素分析工具 V4

特性：
1. ✅ 识别所有类型的元素（input、select、button、label、checkbox、radio等）
2. ✅ 包括隐藏元素（便于了解完整结构）
3. ✅ 详细的定位信息（XPath、ID、Name、Class、data-*属性）
4. ✅ 使用JavaScript深度扫描页面
5. ✅ 记录元素层级关系
6. ✅ 截图每个页面
7. ✅ 生成详细的元素定位Excel

使用方法：
python scripts/analyze_router_pages_v4.py --ip 192.168.3.1 --username admin --password admin1
"""

import sys
import os
import time
import argparse
from datetime import datetime
import json

# 添加项目根目录到 Python 路径
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException, NoSuchElementException, StaleElementReferenceException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from core.router_client import RouterClient


class RouterPageAnalyzerV4:
    """路由器页面完整元素分析器 V4"""

    def __init__(self, router_ip: str, username: str, password: str, model: str = "UR35"):
        self.router_ip = router_ip
        self.username = username
        self.password = password
        self.model = model

        # 创建路由器客户端
        self.router_client = RouterClient(
            router_ip=router_ip,
            username=username,
            password=password,
            model=model
        )

        # 存储分析结果
        self.menu_items = []  # 菜单列表
        self.all_elements = []  # 所有元素列表

        # 创建截图目录
        self.screenshot_dir = r"E:\GIT\ROUTER_TEST\logs\screenshots"
        os.makedirs(self.screenshot_dir, exist_ok=True)

    def analyze(self):
        """执行完整的分析流程"""
        print("\n" + "=" * 80)
        print("路由器页面完整元素分析工具 V4")
        print("=" * 80)
        print(f"路由器IP: {self.router_ip}")
        print(f"用户名: {self.username}")
        print(f"型号: {self.model}")
        print("=" * 80 + "\n")

        try:
            # 1. 登录路由器
            print("步骤1: 登录路由器...")
            if not self.router_client.login_web():
                raise Exception("登录失败")
            print("✅ 登录成功\n")

            # 2. 识别所有菜单项
            print("步骤2: 识别菜单结构...")
            self._collect_all_menu_items()
            print(f"✅ 发现 {len(self.menu_items)} 个页面\n")

            # 3. 遍历每个页面，深度分析所有元素
            if len(self.menu_items) > 0:
                print("步骤3: 深度分析所有页面元素...")
                self._analyze_all_pages_deep()
                print(f"✅ 分析了 {len(self.all_elements)} 个元素\n")
            else:
                print("步骤3: 跳过（没有发现菜单项）\n")

            # 4. 生成详细Excel报告
            print("步骤4: 生成详细Excel报告...")
            output_file = self._generate_detailed_excel()
            print(f"✅ Excel报告已生成: {output_file}\n")

            # 5. 统计信息
            self._print_statistics()

            print("=" * 80)
            print("分析完成！")
            print("=" * 80)

        except Exception as e:
            print(f"\n❌ 分析失败: {e}")
            import traceback
            traceback.print_exc()
        finally:
            # 关闭浏览器
            if self.router_client.driver:
                self.router_client.driver.quit()

    def _collect_all_menu_items(self):
        """收集所有菜单项"""
        driver = self.router_client.driver

        try:
            time.sleep(3)

            try:
                main_menu = driver.find_element(By.ID, "mainmenu")
            except:
                print("  ❌ 未找到主菜单容器 #mainmenu")
                return

            main_menu_items = main_menu.find_elements(By.XPATH, ".//li[contains(@class, 'dropdown')]")
            print(f"  发现 {len(main_menu_items)} 个主菜单项")

            for idx, main_item in enumerate(main_menu_items, 1):
                try:
                    main_link = main_item.find_element(By.XPATH, "./a")
                    main_text = main_link.text.strip()
                    main_target = main_link.get_attribute('data-target')

                    print(f"\n  [{idx}] 主菜单: {main_text}")

                    if main_target:
                        menu_info = {
                            'main_menu': main_text,
                            'sub_menu': '',
                            'text': main_text,
                            'target': main_target
                        }
                        self.menu_items.append(menu_info)
                        print(f"      ✓ {main_text} -> {main_target}")

                    try:
                        if 'dropdown-toggle' in main_link.get_attribute('class'):
                            try:
                                main_link.click()
                                time.sleep(0.5)
                            except:
                                pass

                        sub_menu = main_item.find_element(By.XPATH, ".//ul[contains(@class, 'dropdown-menu')]")
                        sub_items = sub_menu.find_elements(By.XPATH, ".//li/a")

                        print(f"      子菜单数量: {len(sub_items)}")

                        for sub_item in sub_items:
                            try:
                                sub_text = sub_item.text.strip()
                                sub_target = sub_item.get_attribute('data-target')

                                if sub_text and sub_target:
                                    menu_info = {
                                        'main_menu': main_text,
                                        'sub_menu': sub_text,
                                        'text': f"{main_text} > {sub_text}",
                                        'target': sub_target
                                    }
                                    self.menu_items.append(menu_info)
                                    print(f"        ✓ {sub_text} -> {sub_target}")

                            except Exception as e:
                                continue

                    except NoSuchElementException:
                        pass

                except Exception as e:
                    print(f"      ⚠️  主菜单项解析出错: {e}")
                    continue

        except Exception as e:
            print(f"  ❌ 菜单收集失败: {e}")

    def _analyze_all_pages_deep(self):
        """深度分析所有页面"""
        driver = self.router_client.driver
        total = len(self.menu_items)

        for idx, menu_item in enumerate(self.menu_items, 1):
            print(f"\n  [{idx}/{total}] 深度分析: {menu_item['text']}")

            try:
                # 导航到页面
                url = f"http://{self.router_ip}/#{menu_item['target']}"
                driver.get(url)
                time.sleep(3)  # 等待页面和动态内容加载

                # 截图
                safe_filename = menu_item['text'].replace('>', '_').replace('/', '_')[:50]
                screenshot_path = self._take_screenshot(f"page_{idx}_{safe_filename}")

                # 使用JavaScript深度扫描页面
                page_elements = self._deep_scan_page(menu_item, screenshot_path)

                self.all_elements.extend(page_elements)

                print(f"      ✅ 找到 {len(page_elements)} 个元素")

                # 显示元素类型统计
                elem_types = {}
                for elem in page_elements:
                    elem_type = elem['element_type']
                    elem_types[elem_type] = elem_types.get(elem_type, 0) + 1

                print(f"         类型统计: {', '.join([f'{k}({v})' for k, v in elem_types.items()])}")

            except Exception as e:
                print(f"      ⚠️  页面分析出错: {e}")
                continue

    def _deep_scan_page(self, menu_item, screenshot_path):
        """使用JavaScript深度扫描页面所有元素"""
        driver = self.router_client.driver
        elements = []

        try:
            # JavaScript脚本：扫描所有交互元素
            scan_script = """
            function getXPath(element) {
                if (element.id !== '')
                    return '//*[@id="' + element.id + '"]';
                if (element === document.body)
                    return '/html/body';

                var ix = 0;
                var siblings = element.parentNode.childNodes;
                for (var i = 0; i < siblings.length; i++) {
                    var sibling = siblings[i];
                    if (sibling === element) {
                        var tagName = element.tagName.toLowerCase();
                        return getXPath(element.parentNode) + '/' + tagName + '[' + (ix + 1) + ']';
                    }
                    if (sibling.nodeType === 1 && sibling.tagName === element.tagName)
                        ix++;
                }
            }

            function getElementInfo(element) {
                var info = {
                    tag: element.tagName.toLowerCase(),
                    type: element.type || '',
                    id: element.id || '',
                    name: element.name || '',
                    className: element.className || '',
                    value: element.value || '',
                    placeholder: element.placeholder || '',
                    text: element.textContent ? element.textContent.trim().substring(0, 100) : '',
                    href: element.href || '',
                    xpath: getXPath(element),
                    isVisible: element.offsetParent !== null,
                    isDisabled: element.disabled || false,
                    isChecked: element.checked || false,
                    isRequired: element.required || false,
                    dataAttributes: {}
                };

                // 获取所有data-*属性
                for (var i = 0; i < element.attributes.length; i++) {
                    var attr = element.attributes[i];
                    if (attr.name.startsWith('data-')) {
                        info.dataAttributes[attr.name] = attr.value;
                    }
                }

                // 对于select元素，获取选项
                if (element.tagName.toLowerCase() === 'select') {
                    info.options = [];
                    for (var j = 0; j < element.options.length; j++) {
                        info.options.push({
                            value: element.options[j].value,
                            text: element.options[j].text
                        });
                    }
                }

                // 查找关联的label
                if (element.id) {
                    var label = document.querySelector('label[for="' + element.id + '"]');
                    if (label) {
                        info.labelText = label.textContent.trim();
                    }
                }

                // 如果没找到label，尝试查找父元素或前面的label
                if (!info.labelText) {
                    var parent = element.parentElement;
                    if (parent) {
                        var labels = parent.querySelectorAll('label');
                        if (labels.length > 0) {
                            info.labelText = labels[0].textContent.trim().substring(0, 50);
                        }

                        // 查找ys-label类
                        var ysLabel = parent.querySelector('.ys-label');
                        if (ysLabel) {
                            info.labelText = ysLabel.textContent.trim().substring(0, 50);
                        }
                    }
                }

                return info;
            }

            var allElements = [];

            // 1. 扫描所有input元素
            var inputs = document.querySelectorAll('input');
            for (var i = 0; i < inputs.length; i++) {
                var info = getElementInfo(inputs[i]);
                info.elementType = 'input';
                allElements.push(info);
            }

            // 2. 扫描所有select元素
            var selects = document.querySelectorAll('select');
            for (var i = 0; i < selects.length; i++) {
                var info = getElementInfo(selects[i]);
                info.elementType = 'select';
                allElements.push(info);
            }

            // 3. 扫描所有textarea元素
            var textareas = document.querySelectorAll('textarea');
            for (var i = 0; i < textareas.length; i++) {
                var info = getElementInfo(textareas[i]);
                info.elementType = 'textarea';
                allElements.push(info);
            }

            // 4. 扫描所有button元素
            var buttons = document.querySelectorAll('button');
            for (var i = 0; i < buttons.length; i++) {
                var info = getElementInfo(buttons[i]);
                info.elementType = 'button';
                allElements.push(info);
            }

            // 5. 扫描所有label元素
            var labels = document.querySelectorAll('label');
            for (var i = 0; i < labels.length; i++) {
                var info = getElementInfo(labels[i]);
                info.elementType = 'label';
                info.forElement = labels[i].getAttribute('for') || '';
                allElements.push(info);
            }

            // 6. 扫描所有链接
            var links = document.querySelectorAll('a');
            for (var i = 0; i < links.length; i++) {
                if (links[i].getAttribute('data-target') || links[i].onclick) {
                    var info = getElementInfo(links[i]);
                    info.elementType = 'link';
                    allElements.push(info);
                }
            }

            // 7. 扫描带有onclick的div（自定义按钮）
            var divs = document.querySelectorAll('div[onclick]');
            for (var i = 0; i < divs.length; i++) {
                var info = getElementInfo(divs[i]);
                info.elementType = 'div-button';
                allElements.push(info);
            }

            return allElements;
            """

            # 执行JavaScript扫描
            js_elements = driver.execute_script(scan_script)

            # 转换为Python对象
            for idx, js_elem in enumerate(js_elements, 1):
                element_info = {
                    'page_id': idx,
                    'main_menu': menu_item['main_menu'],
                    'sub_menu': menu_item['sub_menu'],
                    'page_target': menu_item['target'],
                    'screenshot': screenshot_path,

                    # 元素基本信息
                    'element_type': js_elem.get('elementType', ''),
                    'tag': js_elem.get('tag', ''),
                    'type': js_elem.get('type', ''),
                    'id': js_elem.get('id', ''),
                    'name': js_elem.get('name', ''),
                    'class': js_elem.get('className', ''),

                    # 元素值
                    'value': js_elem.get('value', ''),
                    'placeholder': js_elem.get('placeholder', ''),
                    'text': js_elem.get('text', ''),
                    'href': js_elem.get('href', ''),

                    # 定位信息
                    'xpath': js_elem.get('xpath', ''),

                    # 状态
                    'is_visible': js_elem.get('isVisible', False),
                    'is_disabled': js_elem.get('isDisabled', False),
                    'is_checked': js_elem.get('isChecked', False),
                    'is_required': js_elem.get('isRequired', False),

                    # 关联信息
                    'label_text': js_elem.get('labelText', ''),
                    'for_element': js_elem.get('forElement', ''),

                    # data属性
                    'data_attributes': json.dumps(js_elem.get('dataAttributes', {})),

                    # select选项
                    'options': json.dumps(js_elem.get('options', [])) if js_elem.get('options') else ''
                }

                elements.append(element_info)

        except Exception as e:
            print(f"        JavaScript扫描出错: {e}")

        return elements

    def _take_screenshot(self, name):
        """截图并返回文件路径"""
        try:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"{name}_{timestamp}.png"
            filepath = os.path.join(self.screenshot_dir, filename)
            self.router_client.driver.save_screenshot(filepath)
            return filename
        except Exception as e:
            return ''

    def _generate_detailed_excel(self):
        """生成详细Excel报告"""
        wb = Workbook()

        if 'Sheet' in wb.sheetnames:
            del wb['Sheet']

        # Sheet1: 菜单导航
        ws_menu = wb.create_sheet("菜单导航", 0)
        self._create_menu_sheet(ws_menu)

        # Sheet2: 所有元素详情
        ws_all = wb.create_sheet("所有元素详情", 1)
        self._create_all_elements_sheet(ws_all)

        # Sheet3: 按类型分类
        ws_by_type = wb.create_sheet("按类型分类", 2)
        self._create_by_type_sheet(ws_by_type)

        # Sheet4: 可配置元素
        ws_config = wb.create_sheet("可配置元素", 3)
        self._create_configurable_sheet(ws_config)

        # 保存文件
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_dir = r"E:\GIT\ROUTER_TEST\config"
        os.makedirs(output_dir, exist_ok=True)
        output_file = os.path.join(output_dir, f"router_full_analysis_v4_{timestamp}.xlsx")
        wb.save(output_file)

        return output_file

    def _create_menu_sheet(self, ws):
        """创建菜单导航Sheet"""
        headers = ['序号', '主菜单', '子菜单', '完整路径', 'data-target', '元素数量']
        ws.append(headers)

        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for idx, item in enumerate(self.menu_items, 1):
            elem_count = len([e for e in self.all_elements if e['page_target'] == item['target']])
            ws.append([
                idx,
                item['main_menu'],
                item['sub_menu'],
                item['text'],
                item['target'],
                elem_count
            ])

        for col, width in [('A', 8), ('B', 20), ('C', 30), ('D', 40), ('E', 40), ('F', 12)]:
            ws.column_dimensions[col].width = width

    def _create_all_elements_sheet(self, ws):
        """创建所有元素详情Sheet"""
        headers = ['序号', '主菜单', '子菜单', 'data-target', '元素类型', 'Tag', 'Type',
                   'ID', 'Name', 'Class', 'Label文本', '默认值', 'Placeholder', '文本内容',
                   'XPath', '是否可见', '是否禁用', '是否必填', '选项列表', 'Data属性', '截图']
        ws.append(headers)

        header_fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        for idx, elem in enumerate(self.all_elements, 1):
            ws.append([
                idx,
                elem['main_menu'],
                elem['sub_menu'],
                elem['page_target'],
                elem['element_type'],
                elem['tag'],
                elem['type'],
                elem['id'],
                elem['name'],
                elem['class'][:50] if elem['class'] else '',
                elem['label_text'][:50] if elem['label_text'] else '',
                elem['value'][:50] if elem['value'] else '',
                elem['placeholder'][:50] if elem['placeholder'] else '',
                elem['text'][:50] if elem['text'] else '',
                elem['xpath'][:100] if elem['xpath'] else '',
                '是' if elem['is_visible'] else '否',
                '是' if elem['is_disabled'] else '否',
                '是' if elem['is_required'] else '否',
                elem['options'][:100] if elem['options'] else '',
                elem['data_attributes'][:100] if elem['data_attributes'] else '',
                elem['screenshot']
            ])

        for col, width in [('A', 8), ('B', 15), ('C', 25), ('D', 30), ('E', 12), ('F', 8),
                          ('G', 10), ('H', 20), ('I', 20), ('J', 30), ('K', 30), ('L', 20),
                          ('M', 20), ('N', 30), ('O', 50), ('P', 10), ('Q', 10), ('R', 10),
                          ('S', 30), ('T', 30), ('U', 40)]:
            ws.column_dimensions[col].width = width

    def _create_by_type_sheet(self, ws):
        """创建按类型分类Sheet"""
        ws.append(['元素类型统计'])
        ws.append([''])

        # 统计各类型数量
        type_counts = {}
        for elem in self.all_elements:
            elem_type = elem['element_type']
            type_counts[elem_type] = type_counts.get(elem_type, 0) + 1

        headers = ['元素类型', '数量', '占比']
        ws.append(headers)

        total = len(self.all_elements)
        for elem_type, count in sorted(type_counts.items(), key=lambda x: x[1], reverse=True):
            percentage = f"{(count/total)*100:.1f}%" if total > 0 else "0%"
            ws.append([elem_type, count, percentage])

        ws.merge_cells('A1:C1')
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

        for col, width in [('A', 20), ('B', 15), ('C', 15)]:
            ws.column_dimensions[col].width = width

    def _create_configurable_sheet(self, ws):
        """创建可配置元素Sheet（仅包含可输入的元素）"""
        ws.append(['路由器配置参数模板（仅可配置项）'])
        ws.append(['说明: 本表只包含可以配置的元素（input、select、textarea等）'])
        ws.append([''])

        headers = ['主菜单', '子菜单', 'data-target', '元素类型', 'Label文本', 'ID', 'Name',
                   'XPath', '默认值', '配置值', '是否启用', '选项列表', '备注']
        ws.append(headers)

        header_fill = PatternFill(start_color='FFC000', end_color='FFC000', fill_type='solid')
        header_font = Font(bold=True, color='000000')
        for cell in ws[4]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # 只包含可配置的元素
        configurable_types = ['input', 'select', 'textarea']
        configurable_elements = [e for e in self.all_elements
                                if e['element_type'] in configurable_types
                                and e['type'] != 'hidden'
                                and e['is_visible']]

        current_page = None
        for elem in configurable_elements:
            page_id = f"{elem['main_menu']}>{elem['sub_menu']}"
            if page_id != current_page:
                if current_page is not None:
                    ws.append([''] * 13)
                current_page = page_id

            ws.append([
                elem['main_menu'],
                elem['sub_menu'],
                elem['page_target'],
                elem['element_type'],
                elem['label_text'],
                elem['id'],
                elem['name'],
                elem['xpath'][:50] if elem['xpath'] else '',
                elem['value'],
                '',  # 配置值
                'Y',  # 是否启用
                elem['options'][:50] if elem['options'] else '',
                ''  # 备注
            ])

        for col, width in [('A', 15), ('B', 25), ('C', 30), ('D', 12), ('E', 25),
                          ('F', 20), ('G', 20), ('H', 40), ('I', 20), ('J', 30),
                          ('K', 10), ('L', 30), ('M', 30)]:
            ws.column_dimensions[col].width = width

        ws.merge_cells('A1:M1')
        ws.merge_cells('A2:M2')
        ws['A1'].font = Font(bold=True, size=14)
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

    def _print_statistics(self):
        """打印统计信息"""
        print("\n" + "=" * 80)
        print("统计信息")
        print("=" * 80)
        print(f"总页面数: {len(self.menu_items)}")
        print(f"总元素数: {len(self.all_elements)}")

        # 按类型统计
        type_counts = {}
        for elem in self.all_elements:
            elem_type = elem['element_type']
            type_counts[elem_type] = type_counts.get(elem_type, 0) + 1

        print("\n元素类型分布:")
        for elem_type, count in sorted(type_counts.items(), key=lambda x: x[1], reverse=True):
            print(f"  {elem_type}: {count}")

        # 可配置元素统计
        configurable = len([e for e in self.all_elements
                           if e['element_type'] in ['input', 'select', 'textarea']
                           and e['type'] != 'hidden'
                           and e['is_visible']])
        print(f"\n可配置元素: {configurable}")
        print("=" * 80)


def main():
    """主函数"""
    parser = argparse.ArgumentParser(description='路由器页面完整元素分析工具 V4')
    parser.add_argument('--ip', required=True, help='路由器IP地址')
    parser.add_argument('--username', default='admin', help='登录用户名（默认: admin）')
    parser.add_argument('--password', default='password', help='登录密码（默认: password）')
    parser.add_argument('--model', default='UR35', help='路由器型号（默认: UR35）')

    args = parser.parse_args()

    analyzer = RouterPageAnalyzerV4(
        router_ip=args.ip,
        username=args.username,
        password=args.password,
        model=args.model
    )

    analyzer.analyze()


if __name__ == '__main__':
    main()
