from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

# 创建文档
doc = Document()

# 设置文档标题
title = doc.add_heading('DMVPN服务器使用指南', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 添加副标题信息
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('版本 1.0 | 2025年11月24日')
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(128, 128, 128)

doc.add_paragraph()

# ==================== 第一部分：服务器信息 ====================
doc.add_heading('一、服务器信息', level=1)

# 创建服务器信息表格
table = doc.add_table(rows=6, cols=2)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

info_data = [
    ('服务器IP', '192.168.50.48'),
    ('操作系统', 'Ubuntu 16.04 LTS'),
    ('登录用户', 'root'),
    ('登录密码', 'milesight123'),
    ('隧道网段', '10.0.0.0/24'),
    ('Hub隧道IP', '10.0.0.1'),
]

for i, (key, value) in enumerate(info_data):
    row = table.rows[i]
    row.cells[0].text = key
    row.cells[1].text = value
    for paragraph in row.cells[0].paragraphs:
        for run in paragraph.runs:
            run.bold = True

doc.add_paragraph()

# ==================== 第二部分：系统架构 ====================
doc.add_heading('二、系统架构', level=1)

doc.add_heading('使用的技术栈', level=2)
tech_list = [
    'IPsec IKE守护进程: racoon (ipsec-tools 0.8.2)',
    '加密协议: IPsec ESP (transport模式)',
    '隧道协议: GRE (Generic Routing Encapsulation)',
    '密钥交换: IKE v1 (Main模式)',
]
for item in tech_list:
    doc.add_paragraph(item, style='List Bullet')

# ==================== 第三部分：服务管理 ====================
doc.add_heading('三、服务管理', level=1)

doc.add_heading('1. 启动服务', level=2)
start_cmds = """# SSH登录服务器
ssh root@192.168.50.48

# 启动racoon服务
systemctl start racoon

# 查看服务状态
systemctl status racoon

# 设置开机自启动
systemctl enable racoon"""
p = doc.add_paragraph()
run = p.add_run(start_cmds)
run.font.name = 'Courier New'
run.font.size = Pt(10)

doc.add_heading('2. 停止服务', level=2)
stop_cmds = """# 停止racoon服务
systemctl stop racoon"""
p = doc.add_paragraph()
run = p.add_run(stop_cmds)
run.font.name = 'Courier New'
run.font.size = Pt(10)

doc.add_heading('3. 重启服务', level=2)
restart_cmds = """# 重启racoon服务
systemctl restart racoon"""
p = doc.add_paragraph()
run = p.add_run(restart_cmds)
run.font.name = 'Courier New'
run.font.size = Pt(10)

# ==================== 第四部分：配置文件 ====================
doc.add_heading('四、配置文件说明', level=1)

doc.add_heading('1. racoon配置文件', level=2)
doc.add_paragraph('文件路径: /etc/racoon/racoon.conf')

racoon_conf = """path pre_shared_key "/etc/racoon/psk.txt";
log notify;

listen {
    isakmp 192.168.50.48 [500];
    isakmp_natt 192.168.50.48 [4500];
}

remote anonymous {
    exchange_mode main;
    passive on;
    generate_policy on;
    nat_traversal on;
    dpd_delay 30;

    proposal {
        encryption_algorithm des;
        hash_algorithm md5;
        authentication_method pre_shared_key;
        dh_group 1;
    }
}

sainfo anonymous {
    encryption_algorithm des;
    authentication_algorithm hmac_md5;
    compression_algorithm deflate;
    lifetime time 3600 sec;
}"""
p = doc.add_paragraph()
run = p.add_run(racoon_conf)
run.font.name = 'Courier New'
run.font.size = Pt(9)

# 配置参数说明表格
doc.add_paragraph()
doc.add_paragraph('配置参数说明:')
param_table = doc.add_table(rows=7, cols=2)
param_table.style = 'Table Grid'

params = [
    ('exchange_mode main', '使用主模式进行IKE协商'),
    ('passive on', 'Hub模式，被动接受连接'),
    ('encryption_algorithm des', 'DES加密算法'),
    ('hash_algorithm md5', 'MD5哈希算法'),
    ('dh_group 1', 'MODP768 Diffie-Hellman组'),
    ('nat_traversal on', '启用NAT穿透'),
    ('dpd_delay 30', 'DPD检测间隔30秒'),
]

for i, (param, desc) in enumerate(params):
    param_table.rows[i].cells[0].text = param
    param_table.rows[i].cells[1].text = desc

doc.add_heading('2. PSK密钥文件', level=2)
doc.add_paragraph('文件路径: /etc/racoon/psk.txt')

psk_content = """192.168.50.16   DMVPNsharedkey123
*               DMVPNsharedkey123"""
p = doc.add_paragraph()
run = p.add_run(psk_content)
run.font.name = 'Courier New'
run.font.size = Pt(10)

doc.add_paragraph('说明: 第一行为特定路由器配置，* 为通配符接受任意IP')

doc.add_heading('3. GRE隧道配置命令', level=2)
gre_cmds = """# 创建GRE隧道
ip tunnel add gre1 mode gre remote 192.168.50.16 local 192.168.50.48 key 123456 ttl 255

# 配置隧道IP
ip addr add 10.0.0.1/32 peer 10.0.0.2 dev gre1

# 启动隧道
ip link set gre1 up"""
p = doc.add_paragraph()
run = p.add_run(gre_cmds)
run.font.name = 'Courier New'
run.font.size = Pt(10)

# ==================== 第五部分：监控和调试 ====================
doc.add_heading('五、监控和调试', level=1)

doc.add_heading('1. 查看IPsec SA状态', level=2)
sa_cmds = """# 查看所有安全关联
setkey -D

# 查看安全策略
setkey -DP"""
p = doc.add_paragraph()
run = p.add_run(sa_cmds)
run.font.name = 'Courier New'
run.font.size = Pt(10)

doc.add_heading('2. 查看日志', level=2)
log_cmds = """# 实时查看日志
tail -f /var/log/syslog | grep racoon

# 查看最近50条日志
tail -50 /var/log/syslog | grep racoon"""
p = doc.add_paragraph()
run = p.add_run(log_cmds)
run.font.name = 'Courier New'
run.font.size = Pt(10)

doc.add_heading('3. 测试连通性', level=2)
test_cmds = """# Ping路由器隧道IP
ping 10.0.0.2

# 检查GRE隧道状态
ip link show gre1

# 检查路由表
ip route show"""
p = doc.add_paragraph()
run = p.add_run(test_cmds)
run.font.name = 'Courier New'
run.font.size = Pt(10)

doc.add_heading('4. 抓包调试', level=2)
tcpdump_cmds = """# 抓取IKE协商流量
tcpdump -i ens33 udp port 500 -n -v

# 抓取ESP加密流量
tcpdump -i ens33 esp -n -v

# 抓取GRE隧道流量
tcpdump -i gre1 -n -v"""
p = doc.add_paragraph()
run = p.add_run(tcpdump_cmds)
run.font.name = 'Courier New'
run.font.size = Pt(10)

# ==================== 第六部分：添加新路由器 ====================
doc.add_heading('六、添加新的Spoke路由器', level=1)

doc.add_heading('路由器端配置要求', level=2)
router_table = doc.add_table(rows=7, cols=2)
router_table.style = 'Table Grid'

router_config = [
    ('Hub地址', '192.168.50.48'),
    ('PSK密钥', 'DMVPNsharedkey123'),
    ('GRE密钥', '123456'),
    ('加密算法', 'DES'),
    ('认证算法', 'MD5'),
    ('DH组', 'MODP768-1'),
    ('PFS组', 'NULL'),
]

for i, (key, value) in enumerate(router_config):
    router_table.rows[i].cells[0].text = key
    router_table.rows[i].cells[1].text = value

# ==================== 第七部分：常见问题 ====================
doc.add_heading('七、常见问题排查', level=1)

doc.add_heading('问题1: 路由器无法连接', level=2)
doc.add_paragraph('错误信息: ERROR: no suitable proposal found')
doc.add_paragraph('解决方法: 检查racoon配置，确保加密参数与路由器匹配（DES/MD5/MODP768）')

doc.add_heading('问题2: Phase 2失败', level=2)
doc.add_paragraph('错误信息: ERROR: pfs group mismatched')
doc.add_paragraph('解决方法: 在sainfo配置中移除pfs_group行，然后重启racoon')

doc.add_heading('问题3: SA建立但GRE不通', level=2)
doc.add_paragraph('排查步骤:')
check_steps = [
    '检查GRE隧道状态: ip link show gre1',
    '检查SPD策略: setkey -DP | grep gre',
    '检查xfrm统计: cat /proc/net/xfrm_stat',
    '抓包验证: tcpdump -i ens33 esp -c 10',
]
for step in check_steps:
    doc.add_paragraph(step, style='List Number')

doc.add_heading('问题4: PSK认证失败', level=2)
doc.add_paragraph('错误信息: ERROR: invalid length of payload')
doc.add_paragraph('解决方法: 检查PSK文件格式，确保没有引号，格式为: IP地址<TAB>密钥')

# ==================== 第八部分：快速参考 ====================
doc.add_heading('八、快速参考', level=1)

doc.add_heading('常用命令速查', level=2)
quick_cmds = """# 服务管理
systemctl start|stop|restart|status racoon

# 查看SA和SPD
setkey -D          # 查看SA
setkey -DP         # 查看SPD

# 刷新配置
setkey -F && setkey -FP && setkey -f /etc/ipsec-tools.conf

# GRE隧道
ip tunnel show gre1
ip addr show gre1

# 测试
ping 10.0.0.2

# 实时日志
tail -f /var/log/syslog | grep racoon"""
p = doc.add_paragraph()
run = p.add_run(quick_cmds)
run.font.name = 'Courier New'
run.font.size = Pt(10)

doc.add_heading('关键端口', level=2)
port_table = doc.add_table(rows=4, cols=2)
port_table.style = 'Table Grid'

ports = [
    ('UDP 500', 'ISAKMP/IKE协商'),
    ('UDP 4500', 'NAT-T (NAT穿透)'),
    ('Protocol 50', 'ESP加密协议'),
    ('Protocol 47', 'GRE隧道协议'),
]

for i, (port, desc) in enumerate(ports):
    port_table.rows[i].cells[0].text = port
    port_table.rows[i].cells[1].text = desc

doc.add_heading('配置文件位置', level=2)
file_table = doc.add_table(rows=4, cols=2)
file_table.style = 'Table Grid'

files = [
    ('/etc/racoon/racoon.conf', 'racoon主配置'),
    ('/etc/racoon/psk.txt', 'PSK密钥文件'),
    ('/etc/ipsec-tools.conf', 'SPD安全策略'),
    ('/var/log/syslog', '系统日志'),
]

for i, (path, desc) in enumerate(files):
    file_table.rows[i].cells[0].text = path
    file_table.rows[i].cells[1].text = desc

# 保存文档到桌面
desktop_path = "C:/Users/admin/Desktop"
output_path = os.path.join(desktop_path, "DMVPN服务器使用指南.docx")
doc.save(output_path)

print(f"文档已保存到: {output_path}")
