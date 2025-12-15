#!/usr/bin/env python
# -*- coding: utf-8 -*-
import openpyxl
import sys

# 读取Excel文件
wb = openpyxl.load_workbook('E:/GIT/ROUTER_TEST/用例步骤.xlsx')
print('所有Sheet:', wb.sheetnames)

# 遍历所有sheet
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f'\n===== Sheet: {sheet_name} =====')
    print(f'行数: {ws.max_row}, 列数: {ws.max_column}')

    # 打印前40行
    print('\n前40行数据:')
    for i in range(1, min(41, ws.max_row + 1)):
        row_data = []
        for j in range(1, ws.max_column + 1):
            cell = ws.cell(row=i, column=j)
            row_data.append(cell.value)
        print(f'Row {i}: {row_data}')

    # 查找ID为24的行
    print('\n查找ID=24的行...')
    for i in range(1, ws.max_row + 1):
        cell_value = ws.cell(row=i, column=1).value
        if cell_value == 24 or cell_value == '24':
            print(f'\n找到ID=24在第{i}行:')
            row_data = []
            for j in range(1, ws.max_column + 1):
                cell = ws.cell(row=i, column=j)
                row_data.append(cell.value)
            print(row_data)
